import os
import re
from copy import copy

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


DEFAULT_VALUES = {
    "*계정": "수입",
    "*과목": "기명후원금",
    "*내역": "기명후원금",
    "생년월일(사업자번호)": "19991231",
    "우편번호": "13809",
    "주 소": "경기도 과천시 홍촌말로 379 (중앙동)",
    "상세주소": "중앙선거관리위원회",
    "직업(업종)": "기타",
    "전화번호": "010-1234-1234",
    "*증빙서첨부": "N",
    "*수입출처구분": "개인",
}

HIGHLIGHT_KEYWORDS = [
    "카카오페이정산",
    "카드자동집금",
    "CMS집금",
    "알림계좌",
    "tosspaymen",
    "엔에이치엔페이",
]

# 연한 붉은색
HIGHLIGHT_FILL = PatternFill(
    fill_type="solid",
    fgColor="FDE9E7"
)


def normalize_header(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"_x[0-9A-Fa-f]{4}_", "", text)

    text = (
        text.replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("\xa0", "")
        .replace(" ", "")
        .strip()
    )

    return text


def clean_amount(value):
    if pd.isna(value):
        return None

    if isinstance(value, (int, float)):
        return int(value)

    value = str(value).replace(",", "").replace("원", "")
    value = re.sub(r"[^\d\-]", "", value)

    if value == "":
        return None

    return int(value)


def clean_date_8digits(value):
    if pd.isna(value):
        return None

    try:
        dt = pd.to_datetime(value)
        return dt.strftime("%Y/%m/%d")
    except:
        pass

    digits = re.sub(r"[^\d]", "", str(value))

    if len(digits) == 8:
        return digits

    return None


def find_header_row(df_raw):
    target_cols = {"거래일자", "입금금액(원)", "거래기록사항"}

    for i in range(min(30, len(df_raw))):
        row_values = set(str(x).strip() for x in df_raw.iloc[i].tolist() if pd.notna(x))
        if len(target_cols.intersection(row_values)) >= 2:
            return i

    raise ValueError("은행내역 파일에서 헤더 행을 찾지 못했습니다.")


def load_bank_data(bank_file_path):
    df_raw = pd.read_excel(bank_file_path, header=None)
    header_row = find_header_row(df_raw)

    df = pd.read_excel(bank_file_path, header=header_row)

    df = df[["거래일자", "입금금액(원)", "거래기록사항"]].copy()
    df = df.dropna(how="all")

    df["*수입일자"] = df["거래일자"].apply(clean_date_8digits)
    df["*수입제공자"] = df["거래기록사항"].astype(str).str.strip()
    df["*금액"] = df["입금금액(원)"].apply(clean_amount)

    df = df.dropna(subset=["*수입일자", "*수입제공자", "*금액"])

    return df[["*수입일자", "*수입제공자", "*금액"]]


def get_template_header_map(ws, header_row=5):
    header_map = {}

    for cell in ws[header_row]:
        if cell.value is not None:
            header_map[normalize_header(cell.value)] = cell.column

    return header_map


def copy_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        source_cell = ws.cell(row=source_row, column=col)
        target_cell = ws.cell(row=target_row, column=col)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        if source_cell.font:
            target_cell.font = copy(source_cell.font)

        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)

        if source_cell.border:
            target_cell.border = copy(source_cell.border)

        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)

        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)

        target_cell.number_format = source_cell.number_format


def find_first_empty_row(ws, start_row=6):
    row = start_row

    while True:
        has_value = any(
            ws.cell(row=row, column=col).value not in [None, ""]
            for col in range(1, ws.max_column + 1)
        )

        if not has_value:
            return row

        row += 1


def should_highlight(provider_name):
    if provider_name is None:
        return False

    text = str(provider_name).strip().lower()

    for keyword in HIGHLIGHT_KEYWORDS:
        if keyword.lower() in text:
            return True

    return False


def highlight_row(ws, row_idx, max_col):
    for col in range(1, max_col + 1):
        ws.cell(row=row_idx, column=col).fill = copy(HIGHLIGHT_FILL)


def write_to_template(template_file_path, output_file_path, df_input):
    wb = load_workbook(template_file_path)
    ws = wb.active

    header_map = get_template_header_map(ws)
    current_row = find_first_empty_row(ws)

    for _, row in df_input.iterrows():
        copy_row_style(ws, 6, current_row, ws.max_column)

        # 기본값 입력
        for header, value in DEFAULT_VALUES.items():
            key = normalize_header(header)

            if key in header_map:
                col = header_map[key]
                ws.cell(row=current_row, column=col, value=value)

        # 은행 데이터 입력
        mapping = {
            "*수입일자": row["*수입일자"],
            "*수입제공자": row["*수입제공자"],
            "*금액": row["*금액"],
        }

        for header, value in mapping.items():
            key = normalize_header(header)

            if key in header_map:
                col = header_map[key]
                ws.cell(row=current_row, column=col, value=value)

        # 특정 키워드 포함 시 행 전체 강조
        if should_highlight(row["*수입제공자"]):
            highlight_row(ws, current_row, max(header_map.values()))

        current_row += 1

    os.makedirs(os.path.dirname(output_file_path), exist_ok=True)
    wb.save(output_file_path)


def run_donation_pipeline(bank_file_path, template_file_path, output_file_path):
    df_input = load_bank_data(bank_file_path)

    write_to_template(
        template_file_path,
        output_file_path,
        df_input,
    )

    return df_input, output_file_path