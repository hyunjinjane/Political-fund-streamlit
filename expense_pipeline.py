import os
import re
import unicodedata
from copy import copy

import pandas as pd
from openpyxl import load_workbook


DEFAULT_VALUES = {
    "*계정": "지출",
    "*과목": "후원금모금경비",
}


def get_target_info_map(config_data):
    targets = config_data.get("targets", [])
    result = {}

    for row in targets:
        target_name = str(row.get("지출대상자", "")).strip()
        if target_name:
            result[target_name] = row

    return result


def get_rule_list(config_data):
    return config_data.get("rules", [])


def get_fallback_config(config_data):
    fallback = config_data.get("fallback", {})
    if not isinstance(fallback, dict):
        fallback = {}
    return fallback


def normalize_header(text):
    if text is None:
        return ""

    text = str(text)
    text = re.sub(r"_x[0-9A-Fa-f]{4}_", "", text)

    return (
        text.replace("\n", "")
        .replace("\r", "")
        .replace("\t", "")
        .replace("\xa0", "")
        .replace(" ", "")
        .strip()
    )


def normalize_text_for_match(text):
    if text is None:
        return ""

    text = str(text)
    text = unicodedata.normalize("NFKC", text)
    text = text.lower().strip()
    text = re.sub(r"[^0-9a-zA-Z가-힣]", "", text)
    return text


def clean_amount(value):
    if pd.isna(value):
        return None

    if isinstance(value, (int, float)):
        return int(value)

    value = str(value).replace(",", "").replace("원", "").strip()
    value = re.sub(r"[^\d\-]", "", value)

    if value == "":
        return None

    try:
        return int(value)
    except Exception:
        return None


def clean_date_8digits(value):
    if pd.isna(value):
        return None

    try:
        dt = pd.to_datetime(value)
        return dt.strftime("%Y/%m/%d")
    except Exception:
        pass

    digits = re.sub(r"[^\d]", "", str(value))
    if len(digits) == 8:
        return f"{digits[:4]}/{digits[4:6]}/{digits[6:]}"

    return None


def find_header_row(df_raw):
    target_cols = {"거래일자", "출금금액(원)", "거래기록사항"}

    for i in range(min(30, len(df_raw))):
        row_values = set(
            normalize_header(x) for x in df_raw.iloc[i].tolist() if pd.notna(x)
        )
        if len(target_cols.intersection(row_values)) >= 2:
            return i

    raise ValueError("은행내역 파일에서 헤더 행을 찾지 못했습니다.")


def standardize_bank_columns(df):
    col_map = {}

    for col in df.columns:
        ncol = normalize_header(col)

        if ncol == "거래일자":
            col_map[col] = "거래일자"
        elif ncol == "출금금액(원)":
            col_map[col] = "출금금액(원)"
        elif ncol == "입금금액(원)":
            col_map[col] = "입금금액(원)"
        elif ncol == "거래기록사항":
            col_map[col] = "거래기록사항"
        elif ncol == "거래내용":
            col_map[col] = "거래내용"

    return df.rename(columns=col_map)


def build_expense_fields(record_text, detail_text="", rules=None):
    if rules is None:
        rules = []

    record = normalize_text_for_match(record_text)
    detail = normalize_text_for_match(detail_text)
    combined = f"{record} {detail}"

    for rule in rules:
        keyword = normalize_text_for_match(rule.get("은행입력키워드", ""))

        if keyword and keyword in combined:
            return {
                "*내역": rule.get("내역", ""),
                "*지출대상자": rule.get("지출대상자", ""),
                "*과목": rule.get("*과목", "후원금모금경비"),
                "*수입지출처구분": rule.get("*수입지출처구분", ""),
                "*증빙서첨부": rule.get("*증빙서첨부", ""),
                "*지출방법": rule.get("*지출방법", ""),
                "force_include": bool(rule.get("금액없어도포함", False)),
            }

    return {
        "*내역": "",
        "*지출대상자": "",
        "*과목": "후원금모금경비",
        "*수입지출처구분": "",
        "*증빙서첨부": "",
        "*지출방법": "",
        "force_include": False,
    }


def load_bank_data(bank_file_path, config_data):
    rules = get_rule_list(config_data)
    fallback = get_fallback_config(config_data)

    df_raw = pd.read_excel(bank_file_path, header=None)
    header_row = find_header_row(df_raw)

    df = pd.read_excel(bank_file_path, header=header_row)
    df = standardize_bank_columns(df)

    required_cols = ["거래일자", "출금금액(원)", "거래기록사항"]
    missing = [col for col in required_cols if col not in df.columns]
    if missing:
        raise ValueError(f"은행내역 파일에 필요한 컬럼이 없습니다: {missing}")

    if "거래내용" not in df.columns:
        df["거래내용"] = ""

    df = df[["거래일자", "출금금액(원)", "거래기록사항", "거래내용"]].copy()
    df = df.dropna(how="all")

    df["*지출일자"] = df["거래일자"].apply(clean_date_8digits)
    df["출금금액_정리"] = df["출금금액(원)"].apply(clean_amount)

    parsed = df.apply(
        lambda row: build_expense_fields(
            record_text=row["거래기록사항"],
            detail_text=row["거래내용"],
            rules=rules
        ),
        axis=1
    ).apply(pd.Series)

    df = pd.concat([df, parsed], axis=1)

    df["include_row"] = (
        df["*지출일자"].notna() &
        (
            df["출금금액_정리"].notna() |
            (df["force_include"] == True)
        )
    )

    df = df[df["include_row"]].copy()

    df["*금액"] = df["출금금액_정리"]

    df["*내역"] = df["*내역"].fillna("")
    df["*지출대상자"] = df["*지출대상자"].fillna("")
    df["*과목"] = df["*과목"].fillna("후원금모금경비")
    df["*수입지출처구분"] = df["*수입지출처구분"].fillna("")
    df["*증빙서첨부"] = df["*증빙서첨부"].fillna("")
    df["*지출방법"] = df["*지출방법"].fillna("")

    mask_unmatched = df["*내역"].astype(str).str.strip().eq("")
    df.loc[mask_unmatched, "*내역"] = df.loc[mask_unmatched, "거래기록사항"].astype(str).str.strip()

    # 규칙 미매칭인데 출금금액이 있는 경우 fallback 전체 적용
    mask_unmatched_with_amount = (
        df["*지출대상자"].astype(str).str.strip().eq("") &
        df["출금금액_정리"].notna()
    )

    if "*과목" in fallback:
        df.loc[mask_unmatched_with_amount, "*과목"] = fallback.get("*과목", "후원금모금경비")

    fallback_detail = str(fallback.get("*내역", "")).strip()
    if fallback_detail:
        df.loc[mask_unmatched_with_amount, "*내역"] = fallback_detail

    df.loc[mask_unmatched_with_amount, "*지출대상자"] = fallback.get("*지출대상자", "")
    df.loc[mask_unmatched_with_amount, "*수입지출처구분"] = fallback.get("*수입지출처구분", "")
    df.loc[mask_unmatched_with_amount, "*증빙서첨부"] = fallback.get("*증빙서첨부", "")
    df.loc[mask_unmatched_with_amount, "*지출방법"] = fallback.get("*지출방법", "")

    # fallback 고정정보 컬럼도 dataframe에 직접 추가
    fallback_fixed_cols = [
        "생년월일(사업자번호)",
        "우편번호",
        "주 소",
        "상세주소",
        "직업(업종)",
        "전화번호",
    ]

    for col_name in fallback_fixed_cols:
        if col_name not in df.columns:
            df[col_name] = ""

        df.loc[mask_unmatched_with_amount, col_name] = fallback.get(col_name, "")

    return df


def get_template_header_map(ws, header_row=5):
    header_map = {}

    for cell in ws[header_row]:
        if cell.value is not None:
            header_map[normalize_header(cell.value)] = cell.column

    return header_map


def copy_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        source_cell = ws.cell(row=source_row, column=col)
        target_cell = ws.cell(row=target_row, column=col)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)

        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)

        target_cell.number_format = source_cell.number_format


def find_first_empty_row(ws, start_row=6):
    row = start_row

    while True:
        has_value = any(
            ws.cell(row=row, column=col).value not in [None, ""]
            for col in range(1, ws.max_column + 1)
        )
        if not has_value:
            return row
        row += 1


def write_to_template(template_file_path, output_file_path, df_input, config_data):
    target_info_map = get_target_info_map(config_data)

    wb = load_workbook(template_file_path)
    ws = wb.active

    header_map = get_template_header_map(ws)
    current_row = find_first_empty_row(ws)

    for _, row in df_input.iterrows():
        copy_row_style(ws, 6, current_row, ws.max_column)

        values_to_write = dict(DEFAULT_VALUES)
        values_to_write["*지출일자"] = row["*지출일자"]
        values_to_write["*내역"] = row["*내역"]
        values_to_write["*지출대상자"] = row["*지출대상자"]
        values_to_write["*과목"] = row["*과목"]
        values_to_write["*수입지출처구분"] = row["*수입지출처구분"]
        values_to_write["*증빙서첨부"] = row["*증빙서첨부"]
        values_to_write["*지출방법"] = row["*지출방법"]

        if pd.notna(row["*금액"]):
            values_to_write["*금액"] = int(row["*금액"])

        target_name = str(row["*지출대상자"]).strip()
        target_info = target_info_map.get(target_name, {})

        fixed_cols = [
            "생년월일(사업자번호)",
            "우편번호",
            "주 소",
            "상세주소",
            "직업(업종)",
            "전화번호",
        ]

        for col_name in fixed_cols:
            row_value = row[col_name] if col_name in row.index else ""

            if pd.notna(row_value) and str(row_value).strip() != "":
                values_to_write[col_name] = row_value
            elif col_name in target_info:
                values_to_write[col_name] = target_info.get(col_name, "")

        for header, value in values_to_write.items():
            key = normalize_header(header)
            if key in header_map:
                col = header_map[key]
                ws.cell(row=current_row, column=col, value=value)

        current_row += 1

    os.makedirs(os.path.dirname(output_file_path), exist_ok=True)
    wb.save(output_file_path)


def run_expense_pipeline(bank_file_path, template_file_path, output_file_path, config_data):
    df_input = load_bank_data(bank_file_path, config_data)
    write_to_template(template_file_path, output_file_path, df_input, config_data)
    return df_input, output_file_path