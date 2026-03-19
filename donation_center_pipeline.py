import os
import json
from copy import copy
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


DEFAULT_CONFIG = {
    "defaults": {
        "주민등록번호(*)": "991231",
        "연락처(*)": "010-1234-1234",
        "직업분류(*)": "기타",
        "직업": "",
        "이메일": ""
    },
    "address": {
        "우편번호(*)": "",
        "주소(*)": "",
        "상세주소(*)": ""
    }
}

REQUIRED_BANK_COLUMNS = [
    "거래기록사항",
    "입금금액(원)",
    "거래일자",
]

EXCLUDE_KEYWORDS = [
    "CMS사용료",
    "카카오페이정산",
    "tosspaymen",
    "엔에리치엔페이",
    "예금이자",
]

HIGHLIGHT_KEYWORDS = [
    "CMS집금",
    "카드자동집금",
    "알림계좌",
]

RED_FILL = PatternFill(fill_type="solid", fgColor="FFCCCC")


def load_json_config(config_source):
    if config_source is None:
        return DEFAULT_CONFIG.copy()

    if isinstance(config_source, dict):
        config = config_source
    else:
        config = json.load(config_source)

    merged = {
        "defaults": DEFAULT_CONFIG["defaults"].copy(),
        "address": DEFAULT_CONFIG["address"].copy()
    }

    merged["defaults"].update(config.get("defaults", {}))
    merged["address"].update(config.get("address", {}))

    return merged


def normalize_date_with_slash(value):
    if pd.isna(value):
        return ""

    try:
        dt = pd.to_datetime(value)
        return dt.strftime("%Y/%m/%d")
    except Exception:
        value = str(value).strip()
        if value.isdigit() and len(value) == 8:
            return f"{value[:4]}/{value[4:6]}/{value[6:8]}"
        return value


def normalize_amount(value):
    if pd.isna(value):
        return None

    if isinstance(value, str):
        value = value.replace(",", "").replace("원", "").strip()
        if value == "":
            return None

    try:
        return int(float(value))
    except Exception:
        return None


def find_header_row(df_raw):
    target_cols = {"거래일자", "입금금액(원)", "거래기록사항"}

    for i in range(min(30, len(df_raw))):
        row_values = set(str(x).strip() for x in df_raw.iloc[i].tolist() if pd.notna(x))
        if len(target_cols.intersection(row_values)) >= 2:
            return i

    raise ValueError("은행내역 파일에서 헤더 행을 찾지 못했습니다.")


def read_uploaded_file(uploaded_file):
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    uploaded_file.seek(0)

    if filename.endswith(".csv"):
        df_raw = pd.read_csv(BytesIO(file_bytes), header=None)
        header_row = find_header_row(df_raw)
        df = pd.read_csv(BytesIO(file_bytes), header=header_row)
    else:
        df_raw = pd.read_excel(BytesIO(file_bytes), header=None)
        header_row = find_header_row(df_raw)
        df = pd.read_excel(BytesIO(file_bytes), header=header_row)

    df.columns = [str(col).strip() for col in df.columns]
    return df


def validate_bank_columns(df):
    missing = [col for col in REQUIRED_BANK_COLUMNS if col not in df.columns]
    if missing:
        raise ValueError(f"은행 파일에 필요한 컬럼이 없습니다: {', '.join(missing)}")


def get_template_headers(ws):
    headers = []
    col_idx = 1

    while True:
        cell_value = ws.cell(row=1, column=col_idx).value
        if cell_value is None:
            break
        headers.append(str(cell_value).strip())
        col_idx += 1

    return headers


def contains_any_keyword(text, keywords):
    if pd.isna(text):
        return False

    text = str(text).strip().lower()
    return any(keyword.lower() in text for keyword in keywords)


def filter_and_mark_bank_df(df_bank):
    df = df_bank.copy()

    df["__normalized_amount"] = df["입금금액(원)"].apply(normalize_amount)
    df["__remark"] = df["거래기록사항"].fillna("").astype(str).str.strip()

    # 하이라이트 대상 여부 먼저 저장
    df["__highlight"] = df["__remark"].apply(
        lambda x: contains_any_keyword(x, HIGHLIGHT_KEYWORDS)
    )

    # 제외 조건
    exclude_mask = (
        df["__normalized_amount"].isna()
        | df["__remark"].apply(lambda x: contains_any_keyword(x, EXCLUDE_KEYWORDS))
    )

    df = df[~exclude_mask].copy()

    return df


def build_center_dataframe(df_bank, user_config, template_headers):
    defaults = user_config.get("defaults", {})
    address_config = user_config.get("address", {})

    result = pd.DataFrame(index=df_bank.index, columns=template_headers)

    for col in template_headers:
        result[col] = ""

    if "성명(*)" in result.columns:
        result["성명(*)"] = df_bank["거래기록사항"].fillna("").astype(str).str.strip()

    if "주민등록번호(*)" in result.columns:
        result["주민등록번호(*)"] = defaults.get("주민등록번호(*)", "991231")

    if "후원금액(*)" in result.columns:
        result["후원금액(*)"] = df_bank["__normalized_amount"]

    if "우편번호(*)" in result.columns:
        result["우편번호(*)"] = address_config.get("우편번호(*)", "")

    if "주소(*)" in result.columns:
        result["주소(*)"] = address_config.get("주소(*)", "")

    if "상세주소(*)" in result.columns:
        result["상세주소(*)"] = address_config.get("상세주소(*)", "")

    converted_dates = df_bank["거래일자"].apply(normalize_date_with_slash)

    if "후원일자(*)" in result.columns:
        result["후원일자(*)"] = converted_dates

    if "입금일자(*)" in result.columns:
        result["입금일자(*)"] = converted_dates

    if "연락처(*)" in result.columns:
        result["연락처(*)"] = defaults.get("연락처(*)", "010-1234-1234")

    if "직업분류(*)" in result.columns:
        result["직업분류(*)"] = defaults.get("직업분류(*)", "기타")

    if "직업" in result.columns:
        result["직업"] = defaults.get("직업", "")

    if "이메일" in result.columns:
        result["이메일"] = defaults.get("이메일", "")

    # 하이라이트 여부 보존
    result["__highlight"] = df_bank["__highlight"].values

    return result


def copy_template_styles(ws, start_row, num_rows, max_col):
    style_row = 2

    for row_idx in range(start_row, start_row + num_rows):
        for col_idx in range(1, max_col + 1):
            source_cell = ws.cell(row=style_row, column=col_idx)
            target_cell = ws.cell(row=row_idx, column=col_idx)

            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)

            if source_cell.font:
                target_cell.font = copy(source_cell.font)

            if source_cell.fill:
                target_cell.fill = copy(source_cell.fill)

            if source_cell.border:
                target_cell.border = copy(source_cell.border)

            if source_cell.alignment:
                target_cell.alignment = copy(source_cell.alignment)

            if source_cell.number_format:
                target_cell.number_format = source_cell.number_format

            if source_cell.protection:
                target_cell.protection = copy(source_cell.protection)


def write_dataframe_to_template(ws, df_result, start_row=2):
    visible_headers = [col for col in df_result.columns if not col.startswith("__")]
    max_col = len(visible_headers)

    copy_template_styles(ws, start_row, len(df_result), max_col)

    for row_offset, (_, row) in enumerate(df_result.iterrows()):
        excel_row = start_row + row_offset

        for col_idx, header in enumerate(visible_headers, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=row[header])

            if bool(row.get("__highlight", False)):
                cell.fill = RED_FILL


def run_center_pipeline(uploaded_file, template_path, output_dir, config_source=None):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"템플릿 파일이 없습니다: {template_path}")

    os.makedirs(output_dir, exist_ok=True)

    df_bank = read_uploaded_file(uploaded_file)
    validate_bank_columns(df_bank)

    df_bank = filter_and_mark_bank_df(df_bank)

    user_config = load_json_config(config_source)

    wb = load_workbook(template_path)
    ws = wb.active

    template_headers = get_template_headers(ws)
    if not template_headers:
        raise ValueError("템플릿 헤더를 읽을 수 없습니다.")

    df_result = build_center_dataframe(df_bank, user_config, template_headers)

    write_dataframe_to_template(ws, df_result, start_row=2)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(output_dir, f"후원회_센터_입력용_{timestamp}.xlsx")
    wb.save(output_path)

    return output_path, df_result[[col for col in df_result.columns if not col.startswith('__')]]
