import os
from io import BytesIO
from copy import copy
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


REQUIRED_BANK_COLUMNS = [
    "거래일자",
    "입금금액(원)",
    "거래기록사항",
]

EXCLUDE_KEYWORDS = [
    "카카오페이정산",
    "tosspaymen",
    "엔에리치엔페이",
    "예금이자",
    "CMS집금",
    "카드자동집금",
    "알림계좌",
]


def normalize_text(value):
    if pd.isna(value):
        return ""
    return str(value).replace(" ", "").strip().lower()


def contains_any_keyword(text, keywords):
    text = normalize_text(text)
    return any(normalize_text(keyword) in text for keyword in keywords)


def normalize_amount(value):
    if pd.isna(value):
        return None

    if isinstance(value, str):
        value = value.replace(",", "").replace("원", "").strip()
        if value == "":
            return None

    try:
        return int(float(value))
    except Exception:
        return None


def normalize_date(value):
    if pd.isna(value):
        return ""

    try:
        dt = pd.to_datetime(value)
        return dt.strftime("%Y-%m-%d")
    except Exception:
        text = str(value).strip()
        if text.isdigit() and len(text) == 8:
            return f"{text[:4]}-{text[4:6]}-{text[6:8]}"
        return text


def find_header_row(df_raw):
    target_cols = {"거래일자", "입금금액(원)", "거래기록사항"}

    for i in range(min(30, len(df_raw))):
        row_values = set(str(x).strip() for x in df_raw.iloc[i].tolist() if pd.notna(x))
        if len(target_cols.intersection(row_values)) >= 2:
            return i

    raise ValueError("은행내역 파일에서 헤더 행을 찾지 못했습니다.")


def read_uploaded_file(uploaded_file):
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()
    uploaded_file.seek(0)

    if filename.endswith(".csv"):
        df_raw = pd.read_csv(BytesIO(file_bytes), header=None)
        header_row = find_header_row(df_raw)
        df = pd.read_csv(BytesIO(file_bytes), header=header_row)
    else:
        df_raw = pd.read_excel(BytesIO(file_bytes), header=None)
        header_row = find_header_row(df_raw)
        df = pd.read_excel(BytesIO(file_bytes), header=header_row)

    df.columns = [str(col).strip() for col in df.columns]
    return df


def validate_bank_columns(df):
    missing = [col for col in REQUIRED_BANK_COLUMNS if col not in df.columns]
    if missing:
        raise ValueError(f"은행 파일에 필요한 컬럼이 없습니다: {', '.join(missing)}")


def filter_bank_df(df_bank):
    df = df_bank.copy()

    df["__amount"] = df["입금금액(원)"].apply(normalize_amount)
    df["__remark"] = df["거래기록사항"].fillna("").astype(str)

    exclude_mask = (
        df["__amount"].isna()
        | df["__remark"].apply(lambda x: contains_any_keyword(x, EXCLUDE_KEYWORDS))
    )

    return df[~exclude_mask].copy().reset_index(drop=True)


def build_report_dataframe(df_bank):
    result = pd.DataFrame()
    result["연번"] = range(1, len(df_bank) + 1)
    result["입금일"] = df_bank["거래일자"].apply(normalize_date)
    result["입금명의"] = df_bank["거래기록사항"].fillna("").astype(str).str.strip()
    result["입금액"] = df_bank["__amount"]
    result["성명"] = ""
    result["연락처"] = ""
    result["비고"] = ""
    return result


def find_report_header_row(ws):
    for row in range(1, min(ws.max_row, 100) + 1):
        values = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
        normalized = [str(v).strip() if v is not None else "" for v in values]

        if "연번" in normalized and "입금일" in normalized and "입금명의" in normalized and "입금액" in normalized:
            return row

    raise ValueError("회보서 템플릿에서 표 헤더 행을 찾지 못했습니다.")


def get_column_map(ws, header_row):
    mapping = {}

    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=col).value
        if value is None:
            continue

        text = str(value).strip()
        if text in ["연번", "입금일", "입금명의", "입금액", "성명", "연락처", "비고"]:
            mapping[text] = col

    required_headers = ["연번", "입금일", "입금명의", "입금액", "성명", "연락처", "비고"]
    missing = [h for h in required_headers if h not in mapping]
    if missing:
        raise ValueError(f"회보서 템플릿 표 컬럼을 찾지 못했습니다: {', '.join(missing)}")

    return mapping


def copy_row_style(ws, source_row, target_row, max_col):
    for col_idx in range(1, max_col + 1):
        source_cell = ws.cell(row=source_row, column=col_idx)
        target_cell = ws.cell(row=target_row, column=col_idx)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.font:
            target_cell.font = copy(source_cell.font)
        if source_cell.fill:
            target_cell.fill = copy(source_cell.fill)
        if source_cell.border:
            target_cell.border = copy(source_cell.border)
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)


def copy_row_height(ws, source_row, target_row):
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def get_horizontal_merge_ranges(ws, row_idx):
    ranges = []
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row == row_idx and merged_range.max_row == row_idx:
            ranges.append((merged_range.min_col, merged_range.max_col))
    return ranges


def copy_horizontal_merges(ws, source_row, target_row):
    for min_col, max_col in get_horizontal_merge_ranges(ws, source_row):
        ws.merge_cells(
            start_row=target_row,
            start_column=min_col,
            end_row=target_row,
            end_column=max_col
        )


def clear_row_values_safe(ws, row_idx, max_col):
    """
    병합셀은 건너뛰고, 실제 값 쓸 수 있는 셀만 비움
    """
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None


def prepare_rows(ws, data_start_row, needed_rows):
    max_col = ws.max_column
    template_row = data_start_row

    if needed_rows > 1:
        ws.insert_rows(data_start_row + 1, needed_rows - 1)

    for row_idx in range(data_start_row, data_start_row + needed_rows):
        if row_idx != template_row:
            copy_row_style(ws, template_row, row_idx, max_col)
            copy_row_height(ws, template_row, row_idx)
            copy_horizontal_merges(ws, template_row, row_idx)

        clear_row_values_safe(ws, row_idx, max_col)


def write_report_data(ws, df_result, data_start_row, column_map):
    for i, (_, row) in enumerate(df_result.iterrows()):
        excel_row = data_start_row + i

        ws.cell(excel_row, column_map["연번"]).value = row["연번"]
        ws.cell(excel_row, column_map["입금일"]).value = row["입금일"]
        ws.cell(excel_row, column_map["입금명의"]).value = row["입금명의"]
        ws.cell(excel_row, column_map["입금액"]).value = row["입금액"]
        ws.cell(excel_row, column_map["성명"]).value = row["성명"]
        ws.cell(excel_row, column_map["연락처"]).value = row["연락처"]
        ws.cell(excel_row, column_map["비고"]).value = row["비고"]


def run_report_pipeline(uploaded_file, template_path, output_dir):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"템플릿 파일이 없습니다: {template_path}")

    os.makedirs(output_dir, exist_ok=True)

    df_bank = read_uploaded_file(uploaded_file)
    validate_bank_columns(df_bank)

    df_bank = filter_bank_df(df_bank)
    df_result = build_report_dataframe(df_bank)

    wb = load_workbook(template_path)
    ws = wb.active

    header_row = find_report_header_row(ws)
    data_start_row = header_row + 1
    column_map = get_column_map(ws, header_row)

    needed_rows = max(len(df_result), 1)
    prepare_rows(ws, data_start_row, needed_rows)

    if len(df_result) > 0:
        write_report_data(ws, df_result, data_start_row, column_map)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = os.path.join(output_dir, f"후원회_회보서_정리_{timestamp}.xlsx")
    wb.save(output_path)

    return output_path, df_result
